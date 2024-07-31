from django.http import JsonResponse
import os
from django.conf import settings
import fitz  # PyMuPDF
from docx import Document
import openpyxl
import pandas as pd
import pytesseract
from PIL import Image
import io
# Assuming extract_text_pdf, extract_text_image, extract_text_docx, and extract_text_xlsx are defined as before

def extract_metadata_pdf(file_path):
    with fitz.open(file_path) as doc:
        metadata = doc.metadata
        return {
            'author': metadata.get('author', 'Unknown'),
            'title': metadata.get('title', 'Unknown'),
            'date': metadata.get('creationDate', 'Unknown'),
            'size': os.path.getsize(file_path)
        }

def extract_metadata_docx(file_path):
    doc = Document(file_path)
    props = doc.core_properties
    return {
        'author': props.author,
        'title': props.title,
        'date': props.created,
        'size': os.path.getsize(file_path)
    }

def extract_metadata_xlsx(file_path):
    wb = openpyxl.load_workbook(file_path)
    props = wb.properties
    return {
        'author': props.creator,
        'title': props.title,
        'date': props.created,
        'size': os.path.getsize(file_path)
    }


def extract_text_pdf(file_path):
    text = ''
    with fitz.open(file_path) as doc:
        for page in doc:
            # Extract text from the page itself
            text += page.get_text()

            # Extract images from the page
            image_list = page.get_images(full=True)
            for image_index, img in enumerate(page.get_images(full=True)):
                # get the XREF of the image
                xref = img[0]
                # extract the image bytes
                base_image = doc.extract_image(xref)
                image_bytes = base_image["image"]

                # you can save the image to disk with open(f"image{page.number}_{image_index}.png", "wb") as f: f.write(image_bytes)
                # or, for OCR, convert it to a PIL Image
                image = Image.open(io.BytesIO(image_bytes))
                # Use pytesseract to do OCR on the image
                image_text = pytesseract.image_to_string(image)
                text += '\n' + image_text
    return text

def extract_text_image(file_path):
    return pytesseract.image_to_string(Image.open(file_path))

def extract_text_docx(file_path):
    doc = Document(file_path)
    return '\n'.join([paragraph.text for paragraph in doc.paragraphs])

def extract_text_xlsx(file_path):
    df = pd.read_excel(file_path)
    return df.to_string()

def index_all_files(request):
    path = os.path.join(settings.STATICFILES_DIRS[0], 'comparative', 'adverts')
    files = os.listdir(path)
    extracted_Json = {}
    for file in files:
        file_path = os.path.join(path, file)
        if file.lower().endswith('.pdf'):
            metadata = extract_metadata_pdf(file_path)
            text = extract_text_pdf(file_path)
        elif file.lower().endswith('.docx'):
            metadata = extract_metadata_docx(file_path)
            text = extract_text_docx(file_path)
        elif file.lower().endswith('.xlsx'):
            metadata = extract_metadata_xlsx(file_path)
            text = extract_text_xlsx(file_path)
        elif file.lower().endswith(('.jpg', '.png', '.jpeg')):
            text = extract_text_image(file_path)
            metadata = {'path': file_path, 'author': 'Unknown', 'title': 'Unknown', 'date': 'Unknown', 'size': os.path.getsize(file_path)}
        elif file.lower().endswith(('.txt', '.json', '.csv')):
            with open(file_path, 'r', encoding='utf-8') as f:
                text = f.read()
            metadata = {'path': file_path, 'author': 'Unknown', 'title': 'Unknown', 'date': 'Unknown', 'size': os.path.getsize(file_path)}
        else:
            extracted_Json[file] = "Unsupported file type."
            continue

        extracted_Json[file] = {**metadata, 'text': text}
    return JsonResponse(extracted_Json)